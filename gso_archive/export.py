"""CSV and Excel exports. Missing values are left blank, never zero."""

from __future__ import annotations

import csv
import io
from datetime import datetime

from .timeutil import to_myt_iso, to_utc_iso, utcnow

COLUMNS = ["series", "obs_time_utc", "obs_time_myt", "value", "unit", "kind", "quality", "dataset_verification", "is_demo"]

ATTRIBUTION = (
    "Source: Grid System Operator (GSO), Malaysia — https://www.gso.org.my/. "
    "Archived copy; GSO's terms of use apply to reuse of this data."
)


def rows_for_export(series_payloads: list[tuple[dict, object]]) -> list[list]:
    """series_payloads: [(points_payload, SeriesInfo)] → flat rows. Gap markers are dropped."""
    rows = []
    for payload, info in series_payloads:
        for p in payload["points"]:
            if p[1] is None:
                continue
            t = datetime.fromisoformat(p[0].replace("Z", "+00:00"))
            rows.append([info.key, to_utc_iso(t), to_myt_iso(t), p[1], info.unit, info.kind, p[2],
                         info.verification, info.is_demo])
    rows.sort(key=lambda r: (r[1], r[0]))
    return rows


def to_csv(rows: list[list]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(COLUMNS)
    writer.writerows(rows)
    return buf.getvalue()


def to_xlsx(rows: list[list], query: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFGHI", (32, 22, 27, 12, 8, 10, 14, 24, 9)):
        ws.column_dimensions[col].width = width

    info = wb.create_sheet("about")
    info.append(["Exported at (UTC)", to_utc_iso(utcnow())])
    for key, value in query.items():
        info.append([key, str(value)])
    info.append(["Rows", len(rows)])
    info.append(["Attribution", ATTRIBUTION])
    info.append(["Notes", "Blank or absent rows mean no reading was archived; they are not zero."])
    if any(r[8] for r in rows):
        info.append(["WARNING", "Contains DEMONSTRATION data — not real GSO readings."])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 100

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
