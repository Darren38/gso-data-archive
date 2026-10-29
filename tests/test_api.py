import io
from datetime import date, datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient

from gso_archive import analytics, db, store
from gso_archive.adapters import Reading, get_adapter
from gso_archive.api.app import create_app
from gso_archive.timeutil import MYT

UTC = timezone.utc
DAY = date(2026, 10, 27)
T0 = datetime(2026, 10, 27, 0, 0, tzinfo=MYT).astimezone(UTC)


def seed(engine, skip_hours=(10, 11)):
    """One MYT day of half-hourly solar, with a gap at 10:00–11:30 MYT."""
    spec = get_adapter("generation_mix")
    rs = []
    for i in range(48):
        t = T0 + timedelta(minutes=30 * i)
        if t.astimezone(MYT).hour in skip_hours:
            continue
        rs.append(Reading("solar", "Solar", t, float(100 * i)))
    with engine.begin() as conn:
        ds = store.ensure_dataset(conn, spec)
        ids = store.resolve_series(conn, spec, ds, rs)
        store.apply_readings(conn, rs, ids, run_id=None, snapshot_id=None, collected_at=datetime.now(UTC))


@pytest.fixture
def api(engine, settings):
    return TestClient(create_app(settings, engine))


def test_empty_archive(api):
    meta = api.get("/api/v1/meta").json()
    assert meta["series"] == [] and meta["is_demo"] is False
    assert any(d.get("no_data_yet") for d in meta["datasets"])
    st = api.get("/api/v1/status").json()
    assert st["totals"]["observations"] == 0 and st["collector_heartbeat"] is None
    assert api.get("/api/v1/health").json()["ok"]


def test_points_mark_gaps_with_null(api, engine):
    seed(engine)
    r = api.get("/api/v1/series/generation_mix.solar/points", params={"start": "2026-10-27", "end": "2026-10-28"})
    assert r.status_code == 200
    pts = r.json()["points"]
    assert len([p for p in pts if p[1] is not None]) == 44
    gaps = [p for p in pts if p[1] is None]
    assert len(gaps) == 1 and gaps[0][2] == "gap"
    assert 0.0 in [p[1] for p in pts]  # a real zero reading stays zero; gaps stay null


def test_daily_stats_and_rollup(api, engine):
    seed(engine)
    rows = api.get("/api/v1/series/generation_mix.solar/daily", params={"start": "2026-10-26", "end": "2026-10-27"}).json()["rows"]
    empty, day = rows
    assert empty["count"] == 0 and empty["max"] is None and empty["energy_mwh"] is None
    assert day["count"] == 44 and day["expected"] == 48
    assert day["max"] == 4700.0 and day["completeness"] == pytest.approx(44 / 48, abs=1e-3)
    # trapezoid across 30-min steps, skipping the 2.5 h gap
    assert 0 < day["energy_mwh"] < sum(100 * i for i in range(48)) * 0.5
    month = api.get("/api/v1/series/generation_mix.solar/daily", params={"start": "2026-10-01", "end": "2026-10-31", "period": "month"}).json()
    assert month["rows"][0]["period"] == "2026-10" and month["rows"][0]["days_with_data"] == 1


def test_csv_export(api, engine):
    seed(engine)
    r = api.get("/api/v1/export", params={"series": "generation_mix.solar", "start": "2026-10-27", "end": "2026-10-27"})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    lines = r.text.strip().splitlines()
    assert lines[0].startswith("series,obs_time_utc,obs_time_myt,value")
    assert len(lines) == 45  # header + 44 readings; the gap is not exported as rows of zero
    assert "2026-10-27T00:00:00+08:00" in lines[1]


def test_xlsx_export(api, engine):
    from openpyxl import load_workbook

    seed(engine)
    r = api.get("/api/v1/export", params={"series": "generation_mix.solar", "start": "2026-10-27", "end": "2026-10-27", "format": "xlsx"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb["data"].max_row == 45
    assert "Attribution" in [c.value for c in wb["about"]["A"]]


@pytest.mark.parametrize("url, code", [
    ("/api/v1/series/bad key/points", 400),
    ("/api/v1/series/generation_mix.nope/points", 404),
    ("/api/v1/series/generation_mix.solar/points?start=garbage", 400),
    ("/api/v1/series/generation_mix.solar/points?start=2020-01-01&end=2026-01-01", 400),
    ("/api/v1/series/generation_mix.solar/daily?period=decade", 422),
    ("/api/v1/export?series=", 400),
    ("/api/v1/export?series=generation_mix.solar&format=pdf", 422),
])
def test_input_validation(api, engine, url, code):
    seed(engine)
    assert api.get(url).status_code == code


def test_no_write_endpoints(api):
    assert api.post("/api/v1/meta").status_code == 405
    assert api.delete("/api/v1/series/generation_mix.solar/points").status_code in (404, 405)


def test_dashboard_served(api):
    r = api.get("/")
    assert r.status_code == 200 and "GSO Grid Archive" in r.text
    assert r.headers["x-content-type-options"] == "nosniff"


def test_status_freshness_and_failures(api, engine):
    seed(engine)
    with engine.begin() as conn:
        rid = store.start_run(conn, "generation_mix", DAY, "schedule")
        store.finish_run(conn, rid, status="failed", error_type="network_error", error_message="timeout")
    st = api.get("/api/v1/status").json()
    gen = next(a for a in st["adapters"] if a["adapter"] == "generation_mix")
    assert gen["consecutive_failures"] == 1 and "timeout" in gen["last_error"]
    fr = next(f for f in st["freshness"] if f["series"] == "generation_mix.solar")
    assert fr["latest_obs_time"].startswith("2026-10-27") and fr["interval_s"] == 1800


def test_rollup_week_grouping():
    days = [{"date": f"2026-10-{d:02d}", "count": 1, "completeness": 1.0, "max": float(d), "max_time": "x", "energy_mwh": 1.0} for d in range(25, 28)]
    weeks = analytics.rollup(days, "week")
    assert [w["period"] for w in weeks] == ["2026-W43", "2026-W44"]
    assert weeks[1]["max"] == 27.0 and weeks[1]["energy_mwh"] == 2.0


def test_observed_interval_from_run_log(engine, settings, api):
    from gso_archive.analytics import observed_interval

    with engine.connect() as conn:
        assert observed_interval(conn) is None  # not enough runs yet
    base = datetime(2026, 10, 27, 4, 0, tzinfo=UTC)
    with engine.begin() as conn:
        for i, gap in enumerate([0, 120, 240, 420, 540, 660]):  # gaps 120,120,180,120,120 s
            rid = store.start_run(conn, "generation_mix", DAY, "schedule")
            store.finish_run(conn, rid, status="success")
            conn.execute(db.collection_runs.update().where(db.collection_runs.c.id == rid)
                         .values(started_at=base + timedelta(seconds=gap)))
        rid = store.start_run(conn, "generation_mix", DAY, "auto_backfill")  # other triggers don't count
    with engine.connect() as conn:
        assert observed_interval(conn) == 120
    c = api.get("/api/v1/meta").json()["collection"]
    assert c["observed_interval_seconds"] == 120 and c["interval_seconds"] == settings.collect_interval_seconds


def test_daily_energy_keeps_the_slice_across_midnight():
    from datetime import date, datetime, timedelta, timezone

    from gso_archive.analytics import _daily_energy
    from gso_archive.timeutil import MYT

    # 100 MW, every 10 minutes from 23:00 to 01:00 MYT: one full hour on each side of midnight
    t0 = datetime(2026, 10, 27, 23, 0, tzinfo=MYT).astimezone(timezone.utc)
    pts = [(t0 + timedelta(minutes=10 * k), 100.0) for k in range(13)]
    e = _daily_energy(pts, 600)
    assert round(e[date(2026, 10, 27)], 6) == 100.0  # before the fix: 83.3 (the 23:50-00:00 slice was lost)
    assert round(e[date(2026, 10, 28)], 6) == 100.0
    # a slice that straddles midnight (23:55 to 00:05, 100 then 200 MW) is split at midnight, 150 MW there
    a = datetime(2026, 10, 27, 23, 55, tzinfo=MYT).astimezone(timezone.utc)
    e = _daily_energy([(a, 100.0), (a + timedelta(minutes=10), 200.0)], 600)
    assert round(e[date(2026, 10, 27)], 6) == round((100 + 150) / 2 * 5 / 60, 6)
    assert round(e[date(2026, 10, 28)], 6) == round((150 + 200) / 2 * 5 / 60, 6)
